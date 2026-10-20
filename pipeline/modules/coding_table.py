"""
IntonaSphere AI - build a corpus coding table from one or more pipeline runs
============================================================================

Turns the intonasphere_results.xlsx of one or more run_all runs into ONE coding table
for a study: your study design on the left (set, pair, gloss, expected contour...),
the pipeline's measures on the right, semitone formulas, and every value that needs a
check in Praat coloured orange with the reason written in a QC column.

    # 1. once per study: a design file to fill in (one row per recording found)
    python -m modules.coding_table --results ~/Desktop/Rue_children/results --make-design ~/Desktop/rue_design.csv

    # 2. after every run: the table
    python -m modules.coding_table \\
        --results Children=~/Desktop/Rue_children/results Shopkeeper=~/Desktop/Rue_shopkeeper/results \\
        --design ~/Desktop/rue_design.csv --corrections ~/Desktop/rue_corrections.csv \\
        --out ~/Desktop/rue_coding_table.xlsx

--results   one or more results files or results folders (the folder must hold
            intonasphere_results.xlsx). Write Label=path to name the speaker or group
            of that run; without a label the speaker is the name of the folder that was
            analysed (e.g. Rue_children).
--design    CSV, one row per recording. The only required column is "id": the clip
            name (07_wh) or its leading number (07). Every other column is copied as it
            is, in your order (e.g. set, comparison, phrase, gloss, expected contour).
            A "speaker" column, if present, wins over the run label.
--design-by question  (PFC) the design file has one row per QUESTION instead: a column
            "question" (Q1, Q2, Q3, S4 or 1, 2, ...) and your columns (expected contour...).
            Each row fills in every speaker's recording of that question; the question of a
            recording comes from its SENTENCE_ID, or from the end of its name (…tw_Q3).
            Make one with --make-design x.csv --design-by question.
Corpus columns  speaker ID, site, country, region group, age, age group, sex, question
            type and sentence are copied from the run's DATASET sheet when they are filled
            (PFC with a speaker table); --no-corpus-info leaves them out. Without a Label=,
            a PFC recording's speaker is its speaker code.
--corrections  CSV of hand values from Praat: id, column, value, note. "column" is a
            header of the output table (e.g. "F0 onset (Hz)"); an empty value blanks the
            cell. Each correction is coloured blue and logged in the Notes and in the
            CORRECTIONS sheet, so the table can always be rebuilt the same way.

What is flagged (orange cell + reason in "QC flags"):
  - an onset (or a max sitting at the onset) 1.8 x the file's median F0 or more: a likely
    octave doubling as voicing starts; any onset, offset or min 1/1.8 of the median or
    less: halving or creak. A high max or offset at the END is not flagged, since a
    question rise can really go that high;
  - the pipeline's own octave-jump flag in the final syllable (creak or a tracking error);
  - onset or minimum within 1.5 semitones of the pitch floor (creak, or a floor set too
    high), maximum within 1.5 semitones of the ceiling (a peak may be cut off);
  - low F0 coverage of the final vowel, a final syllable shorter than 120 ms;
  - a contour class that disagrees with the F_ToBI boundary tone (fall + H%, rise + L%);
  - words taken from speech recognition instead of your .txt.
A flag is a request to look, not a verdict.
"""
from __future__ import annotations

from . import user_path, CSV_ENCODING

import argparse
import datetime as _dt
import math
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd

RESULTS_NAME = "intonasphere_results.xlsx"

DOUBLING_RATIO = 1.8      # value / file median beyond this (or below its inverse) = likely octave error
EDGE_ST = 1.5             # semitones from the pitch floor / ceiling that count as "at the edge"
SHORT_FINAL_S = 0.120     # final syllable shorter than this: check the alignment
LOW_COVERAGE = 0.5        # share of the final vowel with F0 below this: check

# output columns after the design columns: (header, kind). kind: m = measure, f = formula, i = input
MEASURE_COLS: List[Tuple[str, str]] = [
    ("Speaker", "m"), ("Clip file", "m"), ("Pitch range (Hz)", "m"), ("Measure domain", "m"),
    ("Observed contour (pipeline)", "m"), ("Final movement", "m"), ("F_ToBI nuclear", "m"),
    ("F0 onset (Hz)", "m"), ("F0 offset (Hz)", "m"), ("F0 max (Hz)", "m"), ("F0 min (Hz)", "m"),
    ("F0 final syllable (Hz)", "m"), ("Final syllable duration (ms)", "m"),
    ("F0 range (Hz)", "f"), ("F0 range (st)", "f"), ("Offset minus onset (st)", "f"),
    ("QC flags", "m"), ("Matches expected?", "i"), ("Notes", "m"),
]

# columns copied from the run's DATASET sheet (speaker table + PFC sentence), when filled
CORPUS_COLS: List[Tuple[str, str]] = [
    ("SPEAKER_ID", "Speaker ID"), ("SITE_ID", "Site"), ("COUNTRY", "Country"), ("REGION_GROUP", "Region group"),
    ("AGE", "Age"), ("AGE_GROUP", "Age group"), ("SEX", "Sex"), ("QUESTION_TYPE", "Question type"),
    ("SENTENCE_ID", "Sentence ID"), ("SENTENCE_TEXT", "Sentence"),
]

FILL_FLAG = "F4B183"      # orange: check in Praat
FILL_FIXED = "BDD7EE"     # blue: hand-corrected value
FILL_INPUT = "FFF2CC"     # yellow: for you to fill in
FILL_HEAD = "D9D9D9"
FILL_FORMULA = "F2F2F2"


# ---------------------------------------------------------------------------- reading runs

def _results_file(p: Path) -> Path:
    if p.is_dir():
        f = p / RESULTS_NAME
        if not f.exists():
            raise SystemExit(f"No {RESULTS_NAME} in {p}. Point --results at a run_all --out folder.")
        return f
    if not p.exists():
        raise SystemExit(f"Results file not found: {p}")
    return p


def _default_label(f: Path) -> str:
    """Speaker label when none is given: the analysed folder (…/Rue_children/results/x.xlsx -> Rue_children)."""
    d = f.parent
    if d.name.lower() in ("results", "out", "output") and d.parent.name:
        return d.parent.name
    return d.name


def _pitch_range(f: Path) -> str:
    try:
        par = pd.read_excel(f, "PARAMETERS")
    except Exception:
        return ""
    for _, r in par.iterrows():
        if str(r.get("STEP", "")).strip() == "pitch" and str(r.get("PARAMETER", "")).strip() == "range":
            nums = re.findall(r"[\d.]+", str(r.get("VALUE", "")))
            if len(nums) >= 2:
                return f"{float(nums[0]):g}-{float(nums[1]):g}"
            return str(r.get("VALUE", ""))
    return ""


def read_run(spec: str) -> pd.DataFrame:
    """One --results argument -> one row per recording, with the columns the table needs."""
    label = None
    if "=" in spec and not Path(spec).exists():
        label, spec = spec.split("=", 1)
    f = _results_file(user_path(spec))
    explicit = label is not None
    label = label or _default_label(f)
    fd = pd.read_excel(f, "FINAL_DOMAIN")
    try:
        ds = pd.read_excel(f, "DATASET", dtype=str).set_index("TOKEN_ID")
    except Exception:
        ds = pd.DataFrame()
    extra = {}
    for sheet, cols in (("FTOBI_TOKENS", ["NUCLEAR", "QC_WARNINGS"]),
                        ("ALL_TOKENS", ["TEXT_SOURCE"])):
        try:
            s = pd.read_excel(f, sheet)
            extra[sheet] = s.set_index("TOKEN_ID")[[c for c in cols if c in s]]
        except Exception:
            pass
    try:
        warn = pd.read_excel(f, "DATASET_WARNINGS")
    except Exception:
        warn = pd.DataFrame(columns=["AUDIO_FILE", "WARNING"])
    rng = _pitch_range(f)
    rows = []
    for _, r in fd.iterrows():
        tid = str(r["TOKEN_ID"])
        row = r.to_dict()
        row["_LABEL"], row["_RANGE"], row["_FILE"] = label, rng, str(f)
        for k, _h in CORPUS_COLS:
            v = ds.loc[tid, k] if (tid in ds.index and k in ds.columns) else ""
            v = v.iloc[0] if isinstance(v, pd.Series) else v
            row["DS_" + k] = "" if v is None or str(v) == "nan" else str(v)
        # PFC: without a Label=, the speaker is the PFC speaker code, not the folder name
        sp = row["DS_SPEAKER_ID"]
        if not explicit and sp and sp != tid:
            row["_LABEL"] = sp
        for sheet, tab in extra.items():
            if tid in tab.index:
                for c in tab.columns:
                    v = tab.loc[tid, c]
                    row[c if c not in row else f"{sheet}_{c}"] = v.iloc[0] if isinstance(v, pd.Series) else v
        w = warn[warn["AUDIO_FILE"].astype(str) == str(r.get("AUDIO_FILE", ""))]
        row["_WARN"] = "; ".join(w["WARNING"].astype(str)) if len(w) else ""
        rows.append(row)
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------- matching

def _num(v) -> Optional[float]:
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(x) else x


def _match(design_id: str, token_ids: List[str]) -> Optional[str]:
    """07 -> 07_wh; 7 -> 07_wh; 07_wh -> 07_wh. Exact name first, then the part before the first _."""
    d = design_id.strip()
    if d in token_ids:
        return d
    for t in token_ids:
        head = t.split("_", 1)[0]
        if head == d or (head.isdigit() and d.isdigit() and int(head) == int(d)):
            return t
    for t in token_ids:   # 11aab1tw_Q1 must not pick up 11aab1tw_Q10: the prefix has to end at a "_"
        if t.lower().startswith(d.lower() + "_"):
            return t
    return None


def _qkey(v) -> str:
    """Q3, q3, 3 -> '3'; S4 -> '4'; anything else -> itself, upper case."""
    s = str(v).strip().upper()
    m = re.fullmatch(r"[A-Z_]*?(\d+)", s)
    return str(int(m.group(1))) if m else s


def _token_question(r: dict) -> str:
    """Question of a recording: SENTENCE_ID from the run, else the end of its name (…tw_Q3, …tw3)."""
    for k in ("DS_SENTENCE_ID", "SENTENCE_ID"):
        v = r.get(k)
        if v is not None and str(v).strip() not in ("", "nan"):
            return _qkey(v)
    m = re.search(r"(?:_Q|tw|_)(\d+)$", str(r.get("TOKEN_ID", "")), re.I)
    return str(int(m.group(1))) if m else ""


def _st(a: float, b: float) -> float:
    return 12 * math.log2(a / b)


# ---------------------------------------------------------------------------- one row

def measures(r: dict) -> Tuple[Dict[str, object], Dict[str, str], List[str]]:
    """Pipeline row -> (values by output header, {header: reason} for flagged cells, notes)."""
    phrase = _num(r.get("PHRASE_F0_ONSET_HZ")) is not None
    pre = "PHRASE_F0_" if phrase else "F0_"
    onset = _num(r.get(pre + "ONSET_HZ"))
    offset = _num(r.get(pre + "OFFSET_HZ"))
    fmax = _num(r.get(pre + "MAX_HZ")) if phrase else _num(r.get("F0_PEAK_HZ"))
    fmin = _num(r.get(pre + "MIN_HZ")) if phrase else _num(r.get("F0_VALLEY_HZ"))
    contour = str(r.get("CONTOUR_CLASS") or "")
    nuclear = str(r.get("NUCLEAR") or "") if str(r.get("NUCLEAR")) != "nan" else ""
    dur = _num(r.get("FINAL_SYLLABLE_DURATION"))
    label = str(r.get("FINAL_SYLLABLE_LABEL") or "")
    change, slope = _num(r.get("F0_CHANGE_ST")), _num(r.get("FINAL_SLOPE_ST"))
    obs = f"Final syllable [{label}]: {contour}"
    if change is not None:
        obs += f", {change:+.1f} ST"
    if slope is not None:
        obs += f", slope {slope:+.1f} ST/s"
    pch = _num(r.get("PHRASE_F0_CHANGE_ST"))
    if pch is not None:
        obs += f"; phrase {pch:+.1f} ST onset to offset"

    rnd = lambda x: None if x is None else round(x)
    vals = {
        "Speaker": r["_LABEL"], "Clip file": str(r.get("AUDIO_FILE", "")),
        "Pitch range (Hz)": (f"{_num(r.get('PITCH_FLOOR_HZ')):g}-{_num(r.get('PITCH_CEILING_HZ')):g}"
                             if _num(r.get("PITCH_FLOOR_HZ")) and _num(r.get("PITCH_CEILING_HZ")) else r["_RANGE"]), "Measure domain": "whole phrase" if phrase else "final syllable only",
        "Observed contour (pipeline)": obs, "Final movement": contour[:1].upper() + contour[1:],
        "F_ToBI nuclear": nuclear,
        "F0 onset (Hz)": rnd(onset), "F0 offset (Hz)": rnd(offset), "F0 max (Hz)": rnd(fmax),
        "F0 min (Hz)": rnd(fmin), "F0 final syllable (Hz)": rnd(_num(r.get("F0_MEAN_HZ"))),
        "Final syllable duration (ms)": None if dur is None else round(dur * 1000),
    }
    flags: Dict[str, str] = {}

    def flag(col: str, why: str):
        flags[col] = (flags[col] + "; " + why) if col in flags else why

    med = _num(r.get("F0_MEDIAN_HZ_FILE"))
    if med:
        # A high offset or max at the END is often a real question rise, so doubling is only
        # suspected at the onset, or for a max sitting at the onset (the tracker locking onto
        # the 2nd harmonic as voicing starts). Low values anywhere may be halving or creak.
        t_on, t_max = _num(r.get(pre + "ONSET_TIME")), _num(r.get(pre + "MAX_TIME") if phrase else r.get("F0_PEAK_TIME"))
        max_at_onset = t_on is not None and t_max is not None and abs(t_max - t_on) <= 0.1
        high = [("F0 onset (Hz)", onset)] + ([("F0 max (Hz)", fmax)] if max_at_onset else [])
        for col, v in high:
            if v and v / med >= DOUBLING_RATIO:
                flag(col, f"{col.split(' (')[0]} {v:.0f} Hz = {v / med:.2f} x file median ({med:.0f} Hz) at the start of voicing: octave doubling?")
        for col, v in (("F0 onset (Hz)", onset), ("F0 offset (Hz)", offset), ("F0 min (Hz)", fmin)):
            if v and med / v >= DOUBLING_RATIO:
                flag(col, f"{col.split(' (')[0]} {v:.0f} Hz = 1/{med / v:.2f} of file median ({med:.0f} Hz): octave halving or creak?")
    if str(r.get("QC_OCTAVE_JUMP")).lower() == "true":
        j = _num(r.get("MAX_F0_JUMP_ST"))
        flag("F0 min (Hz)", "octave-jump flag in final syllable" + (f" ({j:.1f} ST between frames)" if j else "")
             + ": creak or tracking error; min and range unreliable")
    lo, hi = _num(r.get("PITCH_FLOOR_HZ")), _num(r.get("PITCH_CEILING_HZ"))
    if lo:
        for col, v in (("F0 onset (Hz)", onset), ("F0 min (Hz)", fmin)):
            if v and _st(v, lo) < EDGE_ST:
                flag(col, f"{v:.0f} Hz is within {EDGE_ST} ST of the {lo:.0f} Hz floor: creak, or floor too high?")
    if hi and fmax and _st(hi, fmax) < EDGE_ST:
        flag("F0 max (Hz)", f"{fmax:.0f} Hz is within {EDGE_ST} ST of the {hi:.0f} Hz ceiling: peak cut off?")
    cov = _num(r.get("FINAL_VOWEL_F0_COVERAGE"))
    if str(r.get("QC_FINAL_VOWEL_LOW_F0")).lower() == "true" or (cov is not None and cov < LOW_COVERAGE):
        flag("F0 final syllable (Hz)", f"F0 found on only {cov:.0%} of the final vowel" if cov is not None
             else "little F0 on the final vowel")
    if dur is not None and dur < SHORT_FINAL_S:
        flag("Final syllable duration (ms)", f"final syllable [{label}] only {dur * 1000:.0f} ms: check its boundaries")
    edge = nuclear.split()[-1] if nuclear else ""
    if contour and edge:
        if contour.endswith("fall") and edge.startswith("H%"):
            flag("Final movement", f"contour '{contour}' but F_ToBI boundary {edge}: listen and decide")
        elif contour.endswith("rise") and edge.startswith("L%"):
            flag("Final movement", f"contour '{contour}' but F_ToBI boundary {edge}: listen and decide")
    src = str(r.get("TEXT_SOURCE", ""))
    if src.startswith("asr"):
        flag("Clip file", "words from speech recognition, not your .txt: check the transcript")

    notes = []
    if not phrase:
        notes.append("Whole-phrase measures missing (run_all without --measure both): onset/offset/max/min are final syllable only.")
    qw = str(r.get("QC_WARNINGS", ""))
    if qw and qw != "nan":
        notes.append("F_ToBI: " + qw)
    if r.get("_WARN"):
        notes.append("Warning: " + str(r["_WARN"]))
    return vals, flags, notes


# ---------------------------------------------------------------------------- building

def build(results: List[str], out: Path, design: Optional[Path] = None,
          corrections: Optional[Path] = None, only_design: bool = False,
          design_by: str = "token", corpus_info: bool = True) -> Path:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    runs = pd.concat([read_run(s) for s in results], ignore_index=True)
    dup = runs["TOKEN_ID"].astype(str)[runs["TOKEN_ID"].astype(str).duplicated()].unique().tolist()
    if dup:
        print("Warning: the same recording appears in more than one run; the LAST one is used:", ", ".join(dup))
    runs = runs.drop_duplicates("TOKEN_ID", keep="last")
    by_tid = {str(r["TOKEN_ID"]): r for r in runs.to_dict("records")}
    tids = list(by_tid)

    table, used = [], set()
    if design and design_by == "question":
        # one design row per question (Q1, Q2, ...): it fills in every speaker's recording of it
        dz = pd.read_csv(user_path(design), dtype=str, keep_default_na=False, encoding=CSV_ENCODING)
        qcol = next((c for c in dz.columns if c.strip().lower() in ("question", "sentence_id", "id")), None)
        if qcol is None:
            raise SystemExit("With --design-by question the design file needs a column 'question' (Q1, Q2, ... or 1, 2, ...).")
        dz = dz.rename(columns={qcol: "Question"})
        spk_col = None
        design_cols = ["ID"] + list(dz.columns)
        qrows = {_qkey(q): d for q, d in zip(dz["Question"], dz.to_dict("records"))}
        order = {k: i for i, k in enumerate(qrows)}
        keyed = sorted(tids, key=lambda t: (order.get(_token_question(by_tid[t]), len(order)), t))
        for t in keyed:
            q = _token_question(by_tid[t])
            if q in qrows:
                table.append({"design": {"ID": t, **qrows[q]}, "tid": t})
            elif not only_design:
                table.append({"design": {"ID": t, "Question": q}, "tid": t, "extra": True})
    else:
        if design:
            dz = pd.read_csv(user_path(design), dtype=str, keep_default_na=False, encoding=CSV_ENCODING)
            idcol = next((c for c in dz.columns if c.strip().lower() == "id"), None)
            if idcol is None:
                raise SystemExit("The design file needs a column called 'id'.")
            dz = dz.rename(columns={idcol: "ID"})
        else:
            dz = pd.DataFrame({"ID": tids})
        spk_col = next((c for c in dz.columns if c.strip().lower() == "speaker"), None)
        design_cols = [c for c in dz.columns if c != spk_col]
        for _, d in dz.iterrows():
            t = _match(str(d["ID"]), tids)
            rec = {"design": d.to_dict(), "tid": t}
            if t:
                used.add(t)
            table.append(rec)
        if not only_design:
            for t in tids:
                if t not in used:
                    table.append({"design": {"ID": t}, "tid": t, "extra": True})

    # corpus columns: only those filled for at least one recording (a film clip has none)
    corpus = []
    if corpus_info:
        for k, h in CORPUS_COLS:
            vals_k = [str(by_tid[t].get("DS_" + k, "")) for t in tids]
            if k == "SPEAKER_ID":
                vals_k = [v for v, t in zip(vals_k, tids) if v != t]
            if any(v.strip() for v in vals_k) and h.lower() not in [c.strip().lower() for c in design_cols]:
                corpus.append((k, h))

    wb = Workbook()
    ws = wb.active
    ws.title = "Coding"
    headers = design_cols + [h for _, h in corpus] + [h for h, _ in MEASURE_COLS]
    col = {h: get_column_letter(i + 1) for i, h in enumerate(headers)}
    fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
    bold = Font(name="Arial", bold=True)
    base = Font(name="Arial")
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i, h)
        c.font, c.fill = bold, fill(FILL_HEAD)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    rowmap: Dict[str, int] = {}
    for n, rec in enumerate(table, 2):
        d, t = rec["design"], rec["tid"]
        for h in design_cols:
            ws[f"{col[h]}{n}"] = d.get(h, "")
        rowmap[str(d.get("ID", ""))] = n
        notes = []
        if rec.get("extra"):
            notes.append("Not in the design file." if design_by != "question"
                         else f"Question '{d.get('Question', '')}' is not in the design file.")
        if t is None:
            notes.append("No pipeline result for this id.")
            ws[f"{col['Notes']}{n}"] = " ".join(notes)
            continue
        r = by_tid[t]
        for k, h in corpus:
            v = r.get("DS_" + k, "")
            ws[f"{col[h]}{n}"] = (int(v) if k == "AGE" and str(v).isdigit() else v) if v != t else ""
        vals, flags, more = measures(r)
        if spk_col and str(d.get(spk_col, "")).strip():
            vals["Speaker"] = d[spk_col]
        for h, v in vals.items():
            ws[f"{col[h]}{n}"] = v
        for h, why in flags.items():
            ws[f"{col[h]}{n}"].fill = fill(FILL_FLAG)
            ws[f"{col[h]}{n}"].comment = Comment(why, "IntonaSphere")
        ws[f"{col['QC flags']}{n}"] = " | ".join(f"{h}: {w}" for h, w in flags.items())
        notes.insert(0, f"Pipeline: Praat raw autocorrelation, range {vals['Pitch range (Hz)']} Hz; run file {Path(r['_FILE']).name} ({r['_LABEL']}).")
        ws[f"{col['Notes']}{n}"] = " ".join(notes + more)

    # hand corrections
    log = []
    if corrections:
        cz = pd.read_csv(user_path(corrections), dtype=str, keep_default_na=False, encoding=CSV_ENCODING)
        cz.columns = [c.strip().lower() for c in cz.columns]
        for need in ("id", "column", "value"):
            if need not in cz.columns:
                raise SystemExit(f"The corrections file needs the columns id, column, value (and note). Missing: {need}")
        for _, c in cz.iterrows():
            h = c["column"].strip()
            n = rowmap.get(c["id"].strip())
            if n is None:
                t = _match(c["id"], [str(k) for k in rowmap])
                n = rowmap.get(t) if t else None
            if n is None or h not in col:
                print(f"Correction skipped (unknown id or column): {c['id']} / {h}")
                continue
            cell = ws[f"{col[h]}{n}"]
            old = cell.value
            v = c["value"].strip()
            newv = None if v == "" else (float(v) if re.fullmatch(r"-?\d+(\.\d+)?", v) else v)
            if isinstance(newv, float) and newv.is_integer():
                newv = int(newv)
            cell.value = newv
            cell.fill = fill(FILL_FIXED)
            why = c.get("note", "").strip()
            cell.comment = Comment(f"Hand correction (was {old}). {why}".strip(), "IntonaSphere")
            nc = ws[f"{col['Notes']}{n}"]
            nc.value = (nc.value or "") + f" Hand correction: {h} {old} -> {'blank' if newv is None else newv}" + (f" ({why})." if why else ".")
            log.append({"id": c["id"], "column": h, "pipeline value": old,
                        "hand value": "blank" if newv is None else newv, "note": why})

    # formulas, styles, validation
    last = len(table) + 1
    S, T, Q, R = (col[h] for h in ("F0 max (Hz)", "F0 min (Hz)", "F0 onset (Hz)", "F0 offset (Hz)"))
    for n in range(2, last + 1):
        ws[f"{col['F0 range (Hz)']}{n}"] = f'=IF(AND(ISNUMBER({S}{n}),ISNUMBER({T}{n})),{S}{n}-{T}{n},"")'
        ws[f"{col['F0 range (st)']}{n}"] = f'=IF(AND(ISNUMBER({S}{n}),ISNUMBER({T}{n}),{T}{n}>0),ROUND(12*LOG({S}{n}/{T}{n},2),2),"")'
        ws[f"{col['Offset minus onset (st)']}{n}"] = f'=IF(AND(ISNUMBER({Q}{n}),ISNUMBER({R}{n}),{Q}{n}>0),ROUND(12*LOG({R}{n}/{Q}{n},2),2),"")'
        for h, kind in MEASURE_COLS:
            c = ws[f"{col[h]}{n}"]
            if kind == "f":
                c.fill = fill(FILL_FORMULA)
            elif kind == "i":
                c.fill = fill(FILL_INPUT)
        for c in ws[n]:
            c.font = base
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv_move = DataValidation(type="list", formula1='"Rise,Fall,Flat,Rise-fall,Fall-rise"', allow_blank=True)
    dv_match = DataValidation(type="list", formula1='"Yes,No,Partly"', allow_blank=True)
    ws.add_data_validation(dv_move)
    ws.add_data_validation(dv_match)
    dv_move.add(f"{col['Final movement']}2:{col['Final movement']}{last}")
    dv_match.add(f"{col['Matches expected?']}2:{col['Matches expected?']}{last}")
    widths = {"Observed contour (pipeline)": 42, "QC flags": 50, "Notes": 60, "Clip file": 22}
    for h, L in col.items():
        ws.column_dimensions[L].width = widths.get(h, 14 if h not in design_cols else 22)
    ws.freeze_panes = "B2"

    # legend
    lg = wb.create_sheet("Legend")
    lines = [
        ("How this table was made", f"python -m modules.coding_table, {_dt.date.today().isoformat()}. Rebuild it the same way after every rerun; hand values come back from the corrections file."),
        ("Orange cell", "Check in Praat. The reason is in the cell's comment and in 'QC flags'. A flag is a request to look, not a verdict."),
        ("Blue cell", "Hand-corrected value from the corrections file (old value in the comment and in CORRECTIONS)."),
        ("Yellow cell", "For you to fill in (Matches expected?)."),
        ("Grey cells", "Formulas: range in Hz, range in semitones = 12 x log2(max/min), offset minus onset in semitones (positive = ends higher). Do not type in them."),
        ("Measure domain", "Whole phrase = first to last voiced frame of the phrase (run_all --measure both). Final syllable only = run without --measure both."),
        ("F0 final syllable (Hz)", "Mean F0 of the final syllable."),
        ("Pitch range (Hz)", "Floor-ceiling used for that run (PARAMETERS sheet of the run)."),
        ("Thresholds", f"Octave error: value >= {DOUBLING_RATIO} x or <= 1/{DOUBLING_RATIO} of the file median. Edge: within {EDGE_ST} ST of floor/ceiling. Short final syllable: < {SHORT_FINAL_S * 1000:.0f} ms. Low final-vowel F0: < {LOW_COVERAGE:.0%} coverage."),
    ]
    lg["A1"], lg["A1"].font = "How to read this table", bold
    for i, (a, b) in enumerate(lines, 3):
        lg[f"A{i}"], lg[f"B{i}"] = a, b
        lg[f"A{i}"].font = bold
        lg[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    lg.column_dimensions["A"].width, lg.column_dimensions["B"].width = 26, 110

    src = wb.create_sheet("SOURCES")
    src.append(["Label (speaker)", "Results file", "Pitch range (Hz)", "Recordings"])
    for (lab, f, rg), g in runs.groupby(["_LABEL", "_FILE", "_RANGE"], sort=False):
        src.append([lab, f, rg, ", ".join(g["TOKEN_ID"].astype(str))])
    if log:
        cs = wb.create_sheet("CORRECTIONS")
        cs.append(list(log[0].keys()))
        for e in log:
            cs.append(list(e.values()))
    for sh in [src] + ([wb["CORRECTIONS"]] if log else []):
        for c in sh[1]:
            c.font, c.fill = bold, fill(FILL_HEAD)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Coding table: {out}  ({len(table)} rows, {len(log)} hand corrections)")
    return out


def make_design(results: List[str], out: Path, design_by: str = "token") -> Path:
    runs = pd.concat([read_run(s) for s in results], ignore_index=True).drop_duplicates("TOKEN_ID", keep="last")
    if design_by == "question":
        qs = []
        for r in runs.to_dict("records"):
            q = _token_question(r)
            if q and q not in [x[0] for x in qs]:
                qs.append((q, r.get("DS_SENTENCE_ID") or f"Q{q}", r.get("DS_QUESTION_TYPE", ""), r.get("DS_SENTENCE_TEXT", "")))
        qs.sort(key=lambda x: int(x[0]) if x[0].isdigit() else 999)
        df = pd.DataFrame([{"question": sid, "question type": qt, "sentence": st, "expected contour": "", "notes": ""}
                           for _, sid, qt, st in qs])
        out.parent.mkdir(parents=True, exist_ok=True)
        df.to_csv(out, index=False, encoding=CSV_ENCODING)
        print(f"Design file to fill in: {out}  ({len(df)} questions over {len(runs)} recordings). Keep 'question'.")
        return out
    cols = ["id", "speaker", "set", "comparison", "phrase", "gloss", "pragmatic function", "expected contour"]
    rows = []
    for r in runs.to_dict("records"):
        rows.append({"id": r["TOKEN_ID"], "speaker": r["_LABEL"], "set": "", "comparison": "",
                     "phrase": "", "gloss": "", "pragmatic function": "", "expected contour": ""})
    out.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows, columns=cols).to_csv(out, index=False, encoding=CSV_ENCODING)
    print(f"Design file to fill in: {out}  ({len(rows)} recordings). Add, rename or delete columns freely; keep 'id'.")
    return out


def main():
    ap = argparse.ArgumentParser(description="Build a corpus coding table from run_all results.")
    ap.add_argument("--results", nargs="+", required=True,
                    help="results files or folders; Label=path names the speaker/group of that run")
    ap.add_argument("--design", help="CSV with an 'id' column + your study columns")
    ap.add_argument("--corrections", help="CSV: id, column, value, note (hand values from Praat)")
    ap.add_argument("--out", help="output .xlsx (default: coding_table.xlsx next to the first results file)")
    ap.add_argument("--only-design", action="store_true", help="leave out recordings that are not in the design file")
    ap.add_argument("--make-design", metavar="CSV", help="write a design file to fill in, then stop")
    ap.add_argument("--design-by", choices=["token", "question"], default="token",
                    help="token: one design row per recording (default); question: one row per question "
                         "(Q1, Q2, ...) that fills in every speaker's recording of it (PFC)")
    ap.add_argument("--no-corpus-info", action="store_true",
                    help="leave out the speaker-table columns (site, country, age group, sex, question type...)")
    a = ap.parse_args()
    if a.make_design:
        make_design(a.results, user_path(a.make_design), a.design_by)
        return
    first = a.results[0].split("=", 1)[-1] if "=" in a.results[0] and not Path(a.results[0]).exists() else a.results[0]
    out = user_path(a.out) if a.out else _results_file(user_path(first)).parent / "coding_table.xlsx"
    build(a.results, out, user_path(a.design) if a.design else None,
          user_path(a.corrections) if a.corrections else None, a.only_design,
          a.design_by, not a.no_corpus_info)


if __name__ == "__main__":
    main()
