import os
from datetime import datetime, timedelta

import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches
from fpdf import FPDF
import srt


OUTPUT_DIR = "outputs"
os.makedirs(OUTPUT_DIR, exist_ok=True)


def read_binary_file(path: str) -> bytes:
    with open(path, "rb") as f:
        return f.read()


def safe_text(value) -> str:
    if value is None:
        return ""
    return str(value)



def pdf_safe_text(text, max_token_length: int = 45) -> str:
    """
    Make text safe for FPDF.

    Fixes:
    - unsupported Unicode by replacing characters FPDF cannot render
    - very long unbroken strings that cause:
      "Not enough horizontal space to render a single character"
    """
    if text is None:
        return ""

    text = str(text)

    # Replace problematic Unicode characters with PDF-safe equivalents
    replacements = {
        "–": "-",
        "—": "-",
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
        "→": "->",
        "×": "x",
        "≈": "~",
        "≤": "<=",
        "≥": ">=",
        "ə": "e",
        "ʁ": "R",
        "ɲ": "ny",
        "ʃ": "sh",
        "ʒ": "zh",
        "ɔ": "o",
        "ɛ": "e",
        "ɡ": "g",
        "ɪ": "i",
        "ʊ": "u",
        "ɣ": "gh",
        "ɾ": "r",
        "ː": ":",
        "͡": "",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    # Break very long unbroken tokens
    safe_tokens = []
    for token in text.split(" "):
        if len(token) > max_token_length:
            chunks = [
                token[i:i + max_token_length]
                for i in range(0, len(token), max_token_length)
            ]
            safe_tokens.append(" ".join(chunks))
        else:
            safe_tokens.append(token)

    text = " ".join(safe_tokens)

    # FPDF default fonts are latin-1 only
    return text.encode("latin-1", "replace").decode("latin-1")


def timestamped_filename(prefix: str, speaker_id: str, ext: str) -> str:
    now = datetime.now().strftime("%Y%m%d_%H%M%S")
    clean_speaker = "".join(
        c for c in speaker_id if c.isalnum() or c in ("-", "_")
    ) or "speaker"
    return f"{clean_speaker}_{prefix}_{now}.{ext}"


# ============================================================
# Basic transcript exports
# ============================================================

def save_transcript_txt(transcript: str, filename: str = "transcript.txt") -> str:
    path = os.path.join(OUTPUT_DIR, filename)
    with open(path, "w", encoding="utf-8") as f:
        f.write(transcript or "")
    return path


def save_transcript_docx(transcript: str, filename: str = "transcript.docx") -> str:
    path = os.path.join(OUTPUT_DIR, filename)

    doc = Document()
    doc.add_heading("Transcript", level=1)
    doc.add_paragraph(transcript or "")
    doc.save(path)

    return path



def pdf_safe_text(text, max_token_length: int = 30) -> str:
    """
    Make text safe for FPDF default fonts and line wrapping.
    """
    if text is None:
        return ""

    text = str(text)

    replacements = {
        "–": "-",
        "—": "-",
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
        "→": "->",
        "×": "x",
        "≈": "~",
        "≤": "<=",
        "≥": ">=",
        "ə": "e",
        "ʁ": "R",
        "ɲ": "ny",
        "ʃ": "sh",
        "ʒ": "zh",
        "ɔ": "o",
        "ɛ": "e",
        "ɡ": "g",
        "ɪ": "i",
        "ʊ": "u",
        "ɣ": "gh",
        "ɾ": "r",
        "ː": ":",
        "͡": "",
        "̃": "",
        "\u0303": "",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    text = text.encode("latin-1", "replace").decode("latin-1")

    # Break long tokens so FPDF can wrap them
    tokens = []
    for token in text.split(" "):
        if len(token) > max_token_length:
            chunks = [
                token[i:i + max_token_length]
                for i in range(0, len(token), max_token_length)
            ]
            tokens.append(" ".join(chunks))
        else:
            tokens.append(token)

    return " ".join(tokens)


def pdf_write_block(pdf, text, line_height=6):
    """
    Write text safely to PDF. This avoids FPDF width errors by:
    - resetting X position before every block
    - using a fixed safe width
    - splitting very long lines
    """
    safe = pdf_safe_text(text)

    page_width = pdf.w - pdf.l_margin - pdf.r_margin

    for raw_line in safe.splitlines() or [""]:
        line = raw_line.strip()

        if not line:
            pdf.ln(line_height)
            continue

        # Additional hard split for very long lines
        while len(line) > 110:
            chunk = line[:110]
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(page_width, line_height, chunk)
            line = line[110:]

        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(page_width, line_height, line)


def save_transcript_pdf(transcript: str, filename: str = "transcript.pdf") -> str:
    path = os.path.join(OUTPUT_DIR, filename)

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    pdf.set_font("Arial", "B", 16)
    pdf_write_block(pdf, "Transcript", line_height=8)
    pdf.ln(3)

    pdf.set_font("Arial", size=11)
    pdf_write_block(pdf, transcript or "", line_height=6)

    pdf.output(path)

    return path


def save_srt(segments: list, filename: str = "subtitles.srt") -> str:
    path = os.path.join(OUTPUT_DIR, filename)

    subtitles = []

    for index, segment in enumerate(segments or [], start=1):
        subtitles.append(
            srt.Subtitle(
                index=index,
                start=timedelta(seconds=float(segment["start"])),
                end=timedelta(seconds=float(segment["end"])),
                content=segment["text"],
            )
        )

    with open(path, "w", encoding="utf-8") as f:
        f.write(srt.compose(subtitles))

    return path


# ============================================================
# Acoustic table exports
# ============================================================

def save_acoustic_csv(df: pd.DataFrame, filename: str = "acoustic_analysis.csv") -> str:
    path = os.path.join(OUTPUT_DIR, filename)

    if df is not None:
        df.to_csv(path, index=False)
    else:
        pd.DataFrame().to_csv(path, index=False)

    return path


def save_acoustic_excel(df: pd.DataFrame, filename: str = "acoustic_analysis.xlsx") -> str:
    path = os.path.join(OUTPUT_DIR, filename)

    if df is not None:
        df.to_excel(path, index=False)
    else:
        pd.DataFrame().to_excel(path, index=False)

    return path


# ============================================================
# Pitch/F0 contour image
# ============================================================

def save_pitch_contour_image(
    pitch_df: pd.DataFrame,
    speaker_id: str = "speaker",
    filename: str = None,
) -> str:
    """
    Save pitch/F0 contour as a PNG image.
    """
    if filename is None:
        filename = f"{speaker_id}_pitch_contour.png"

    path = os.path.join(OUTPUT_DIR, filename)

    fig, ax = plt.subplots(figsize=(10, 5))
    ax.plot(pitch_df["time_seconds"], pitch_df["f0_hz"])
    ax.set_xlabel("Time (seconds)")
    ax.set_ylabel("F0 (Hz)")
    ax.set_title(f"Pitch / F0 Contour - {speaker_id}")
    fig.tight_layout()
    fig.savefig(path, dpi=300)
    plt.close(fig)

    return path


# ============================================================
# Report-building helpers
# ============================================================

def dataframe_to_key_value_lines(df: pd.DataFrame) -> list:
    """
    Convert a one-row dataframe into key-value pairs.
    """
    if df is None or df.empty:
        return []

    row = df.iloc[0].to_dict()
    return [(str(key), safe_text(value)) for key, value in row.items()]


def build_full_analysis_sections(
    speaker_id: str,
    file_name: str,
    orthographic_transcript: str,
    phonetic_transcription: str,
    acoustic_df: pd.DataFrame = None,
    pitch_image_path: str = "",
    accuracy_info: dict = None,
) -> dict:
    """
    Organize all analysis outputs into coherent report sections.
    """
    accuracy_info = accuracy_info or {}

    return {
        "metadata": {
            "Speaker ID": speaker_id,
            "File name": file_name,
            "Report generated": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        },
        "transcription": {
            "Orthographic transcription": orthographic_transcript or "",
            "Phonetic transcription": phonetic_transcription or "",
        },
        "accuracy": accuracy_info,
        "acoustic_lines": dataframe_to_key_value_lines(acoustic_df),
        "pitch_image_path": pitch_image_path,
    }


# ============================================================
# Full coherent analysis exports
# ============================================================

def save_full_analysis_txt(
    speaker_id: str,
    file_name: str,
    orthographic_transcript: str,
    phonetic_transcription: str,
    acoustic_df: pd.DataFrame = None,
    pitch_image_path: str = "",
    accuracy_info: dict = None,
    filename: str = None,
) -> str:
    """
    Save coherent full analysis as TXT.
    """
    if filename is None:
        filename = timestamped_filename("full_analysis_report", speaker_id, "txt")

    path = os.path.join(OUTPUT_DIR, filename)

    sections = build_full_analysis_sections(
        speaker_id=speaker_id,
        file_name=file_name,
        orthographic_transcript=orthographic_transcript,
        phonetic_transcription=phonetic_transcription,
        acoustic_df=acoustic_df,
        pitch_image_path=pitch_image_path,
        accuracy_info=accuracy_info,
    )

    with open(path, "w", encoding="utf-8") as f:
        f.write("JamiSpeak Transcriber + Praat Analyzer\n")
        f.write("Full Analysis Report\n")
        f.write("=" * 72 + "\n\n")

        f.write("1. Speaker and File Information\n")
        f.write("-" * 72 + "\n")
        for key, value in sections["metadata"].items():
            f.write(f"{key}: {value}\n")
        f.write("\n")

        f.write("2. Orthographic Transcription\n")
        f.write("-" * 72 + "\n")
        f.write(sections["transcription"]["Orthographic transcription"] + "\n\n")

        f.write("3. Phonetic Transcription\n")
        f.write("-" * 72 + "\n")
        f.write(sections["transcription"]["Phonetic transcription"] + "\n\n")

        f.write("4. Transcription Accuracy Indices\n")
        f.write("-" * 72 + "\n")
        if sections["accuracy"]:
            for key, value in sections["accuracy"].items():
                f.write(f"{key}: {value}\n")
        else:
            f.write("No accuracy result available.\n")
        f.write("\n")

        f.write("5. Acoustic, Prosodic, TextGrid, and Prominence Analysis\n")
        f.write("-" * 72 + "\n")
        if sections["acoustic_lines"]:
            for key, value in sections["acoustic_lines"]:
                f.write(f"{key}: {value}\n")
        else:
            f.write("No acoustic/prosodic analysis available.\n")
        f.write("\n")

        f.write("6. Pitch/F0 Contour Image\n")
        f.write("-" * 72 + "\n")
        f.write(f"Pitch contour image path: {pitch_image_path or 'Not available'}\n")

    return path


def save_full_analysis_docx(
    speaker_id: str,
    file_name: str,
    orthographic_transcript: str,
    phonetic_transcription: str,
    acoustic_df: pd.DataFrame = None,
    pitch_image_path: str = "",
    accuracy_info: dict = None,
    filename: str = None,
) -> str:
    """
    Save coherent full analysis as DOCX.
    DOCX is best for IPA preservation.
    """
    if filename is None:
        filename = timestamped_filename("full_analysis_report", speaker_id, "docx")

    path = os.path.join(OUTPUT_DIR, filename)

    sections = build_full_analysis_sections(
        speaker_id=speaker_id,
        file_name=file_name,
        orthographic_transcript=orthographic_transcript,
        phonetic_transcription=phonetic_transcription,
        acoustic_df=acoustic_df,
        pitch_image_path=pitch_image_path,
        accuracy_info=accuracy_info,
    )

    doc = Document()

    doc.add_heading("JamiSpeak Transcriber + Praat Analyzer", level=0)
    doc.add_heading("Full Analysis Report", level=1)

    doc.add_heading("1. Speaker and File Information", level=2)
    for key, value in sections["metadata"].items():
        doc.add_paragraph(f"{key}: {value}")

    doc.add_heading("2. Orthographic Transcription", level=2)
    doc.add_paragraph(sections["transcription"]["Orthographic transcription"])

    doc.add_heading("3. Phonetic Transcription", level=2)
    doc.add_paragraph(sections["transcription"]["Phonetic transcription"])

    doc.add_heading("4. Transcription Accuracy Indices", level=2)
    if sections["accuracy"]:
        for key, value in sections["accuracy"].items():
            doc.add_paragraph(f"{key}: {value}")
    else:
        doc.add_paragraph("No accuracy result available.")

    doc.add_heading("5. Acoustic, Prosodic, TextGrid, and Prominence Analysis", level=2)

    if acoustic_df is not None and not acoustic_df.empty:
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"

        hdr = table.rows[0].cells
        hdr[0].text = "Measure"
        hdr[1].text = "Value"

        for key, value in sections["acoustic_lines"]:
            cells = table.add_row().cells
            cells[0].text = key
            cells[1].text = value
    else:
        doc.add_paragraph("No acoustic/prosodic analysis available.")

    doc.add_heading("6. Pitch/F0 Contour", level=2)

    if pitch_image_path and os.path.exists(pitch_image_path):
        doc.add_picture(pitch_image_path, width=Inches(6.5))
    else:
        doc.add_paragraph("No pitch contour image available.")

    doc.save(path)

    return path



def save_full_analysis_pdf(
    speaker_id: str,
    file_name: str,
    orthographic_transcript: str,
    phonetic_transcription: str,
    acoustic_df: pd.DataFrame = None,
    pitch_image_path: str = "",
    accuracy_info: dict = None,
    filename: str = None,
) -> str:
    """
    Save coherent full analysis as PDF.

    Note:
    PDF output is simplified for safety. DOCX is the best format for IPA.
    """
    if filename is None:
        filename = timestamped_filename("full_analysis_report", speaker_id, "pdf")

    path = os.path.join(OUTPUT_DIR, filename)

    sections = build_full_analysis_sections(
        speaker_id=speaker_id,
        file_name=file_name,
        orthographic_transcript=orthographic_transcript,
        phonetic_transcription=phonetic_transcription,
        acoustic_df=acoustic_df,
        pitch_image_path=pitch_image_path,
        accuracy_info=accuracy_info,
    )

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Arial", "B", 16)
    pdf_write_block(pdf, "JamiSpeak Transcriber + Praat Analyzer", line_height=8)
    pdf_write_block(pdf, "Full Analysis Report", line_height=8)
    pdf.ln(3)

    pdf.set_font("Arial", "B", 13)
    pdf_write_block(pdf, "1. Speaker and File Information", line_height=7)
    pdf.set_font("Arial", size=10)

    for key, value in sections["metadata"].items():
        pdf_write_block(pdf, f"{key}: {value}", line_height=6)

    pdf.ln(2)

    pdf.set_font("Arial", "B", 13)
    pdf_write_block(pdf, "2. Orthographic Transcription", line_height=7)
    pdf.set_font("Arial", size=10)
    pdf_write_block(
        pdf,
        sections["transcription"]["Orthographic transcription"],
        line_height=6,
    )

    pdf.ln(2)

    pdf.set_font("Arial", "B", 13)
    pdf_write_block(pdf, "3. Phonetic Transcription", line_height=7)
    pdf.set_font("Arial", size=10)
    pdf_write_block(
        pdf,
        sections["transcription"]["Phonetic transcription"],
        line_height=6,
    )

    pdf.ln(2)

    pdf.set_font("Arial", "B", 13)
    pdf_write_block(pdf, "4. Transcription Accuracy Indices", line_height=7)
    pdf.set_font("Arial", size=10)

    if sections["accuracy"]:
        for key, value in sections["accuracy"].items():
            pdf_write_block(pdf, f"{key}: {value}", line_height=6)
    else:
        pdf_write_block(pdf, "No accuracy result available.", line_height=6)

    pdf.ln(2)

    pdf.set_font("Arial", "B", 13)
    pdf_write_block(
        pdf,
        "5. Acoustic, Prosodic, TextGrid, and Prominence Analysis",
        line_height=7,
    )

    pdf.set_font("Arial", size=8)

    if sections["acoustic_lines"]:
        for key, value in sections["acoustic_lines"]:
            pdf_write_block(pdf, f"{key}: {value}", line_height=5)
    else:
        pdf_write_block(pdf, "No acoustic/prosodic analysis available.", line_height=5)

    pdf.ln(2)

    pdf.set_font("Arial", "B", 13)
    pdf_write_block(pdf, "6. Pitch/F0 Contour", line_height=7)

    if pitch_image_path and os.path.exists(pitch_image_path):
        try:
            pdf.image(pitch_image_path, x=10, w=180)
        except Exception:
            pdf.set_font("Arial", size=10)
            pdf_write_block(pdf, f"Pitch contour image: {pitch_image_path}", line_height=6)
    else:
        pdf.set_font("Arial", size=10)
        pdf_write_block(pdf, "No pitch contour image available.", line_height=6)

    pdf.output(path)

    return path


def save_full_analysis_excel(
    speaker_id: str,
    file_name: str,
    orthographic_transcript: str,
    phonetic_transcription: str,
    acoustic_df: pd.DataFrame = None,
    accuracy_info: dict = None,
    filename: str = None,
) -> str:
    """
    Save coherent full analysis as Excel with multiple sheets.
    """
    if filename is None:
        filename = timestamped_filename("full_analysis_report", speaker_id, "xlsx")

    path = os.path.join(OUTPUT_DIR, filename)

    metadata_df = pd.DataFrame(
        [
            {
                "speaker_id": speaker_id,
                "file_name": file_name,
                "report_generated": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            }
        ]
    )

    transcript_df = pd.DataFrame(
        [
            {
                "orthographic_transcription": orthographic_transcript or "",
                "phonetic_transcription": phonetic_transcription or "",
            }
        ]
    )

    accuracy_df = pd.DataFrame([accuracy_info or {}])

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        metadata_df.to_excel(writer, sheet_name="metadata", index=False)
        transcript_df.to_excel(writer, sheet_name="transcription", index=False)
        accuracy_df.to_excel(writer, sheet_name="accuracy", index=False)

        if acoustic_df is not None and not acoustic_df.empty:
            acoustic_df.to_excel(writer, sheet_name="analysis", index=False)
        else:
            pd.DataFrame().to_excel(writer, sheet_name="analysis", index=False)

    return path


def save_full_analysis_csv(
    acoustic_df: pd.DataFrame,
    speaker_id: str,
    filename: str = None,
) -> str:
    """
    Save the complete analysis dataframe as CSV.
    """
    if filename is None:
        filename = timestamped_filename("full_analysis_report", speaker_id, "csv")

    path = os.path.join(OUTPUT_DIR, filename)

    if acoustic_df is not None:
        acoustic_df.to_csv(path, index=False)
    else:
        pd.DataFrame().to_csv(path, index=False)

    return path


def save_transcription_report_txt(
    speaker_id: str,
    file_name: str,
    orthographic_transcript: str,
    phonetic_transcription: str,
    highest_f0_info: dict = None,
    filename: str = None,
) -> str:
    """
    Short backward-compatible TXT report.
    """
    if filename is None:
        filename = f"{speaker_id}_transcription_report.txt"

    path = os.path.join(OUTPUT_DIR, filename)

    highest_f0_info = highest_f0_info or {}

    with open(path, "w", encoding="utf-8") as f:
        f.write("JamiSpeak Transcriber + Praat Analyzer Report\n")
        f.write("=" * 55 + "\n\n")

        f.write(f"Speaker ID: {speaker_id}\n")
        f.write(f"File name: {file_name}\n\n")

        f.write("Orthographic Transcription\n")
        f.write("-" * 30 + "\n")
        f.write((orthographic_transcript or "").strip() + "\n\n")

        f.write("Phonetic Transcription\n")
        f.write("-" * 30 + "\n")
        f.write((phonetic_transcription or "").strip() + "\n\n")

        f.write("Highest F0 Sound / Interval\n")
        f.write("-" * 30 + "\n")
        if highest_f0_info:
            for key, value in highest_f0_info.items():
                f.write(f"{key}: {value}\n")
        else:
            f.write("No highest-F0 interval information available.\n")

    return path



# ============================================================
# Batch analysis exports
# ============================================================

def save_batch_analysis_csv(
    batch_df: pd.DataFrame,
    filename: str = "batch_acoustic_analysis.csv",
) -> str:
    """
    Save a combined batch dataframe as CSV.
    Each row should represent one speaker/file.
    """
    path = os.path.join(OUTPUT_DIR, filename)

    if batch_df is not None and not batch_df.empty:
        batch_df.to_csv(path, index=False)
    else:
        pd.DataFrame().to_csv(path, index=False)

    return path


def save_batch_analysis_excel(
    batch_df: pd.DataFrame,
    filename: str = "batch_acoustic_analysis.xlsx",
) -> str:
    """
    Save a combined batch dataframe as Excel.
    """
    path = os.path.join(OUTPUT_DIR, filename)

    if batch_df is not None and not batch_df.empty:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            batch_df.to_excel(writer, sheet_name="batch_analysis", index=False)
    else:
        pd.DataFrame().to_excel(path, index=False)

    return path


def save_batch_analysis_txt(
    batch_df: pd.DataFrame,
    filename: str = "batch_acoustic_analysis.txt",
) -> str:
    """
    Save a combined batch dataframe as a readable TXT file.
    """
    path = os.path.join(OUTPUT_DIR, filename)

    with open(path, "w", encoding="utf-8") as f:
        f.write("JamiSpeak Batch Acoustic / Prosodic Analysis\n")
        f.write("=" * 80 + "\n\n")

        if batch_df is None or batch_df.empty:
            f.write("No batch analysis data available.\n")
            return path

        f.write(f"Number of rows/files: {len(batch_df)}\n")
        f.write(f"Number of variables: {len(batch_df.columns)}\n\n")

        for idx, row in batch_df.iterrows():
            speaker_id = row.get("speaker_id", f"row_{idx + 1}")
            file_name = row.get("file_name", "")

            f.write("-" * 80 + "\n")
            f.write(f"Row: {idx + 1}\n")
            f.write(f"Speaker ID: {speaker_id}\n")
            f.write(f"File name: {file_name}\n")
            f.write("-" * 80 + "\n")

            for col in batch_df.columns:
                value = row.get(col, "")
                f.write(f"{col}: {value}\n")

            f.write("\n")

    return path


def save_batch_analysis_docx(
    batch_df: pd.DataFrame,
    filename: str = "batch_acoustic_analysis.docx",
) -> str:
    """
    Save a combined batch dataframe as DOCX.
    """
    path = os.path.join(OUTPUT_DIR, filename)

    doc = Document()
    doc.add_heading("JamiSpeak Batch Acoustic / Prosodic Analysis", level=0)

    if batch_df is None or batch_df.empty:
        doc.add_paragraph("No batch analysis data available.")
        doc.save(path)
        return path

    doc.add_paragraph(f"Number of rows/files: {len(batch_df)}")
    doc.add_paragraph(f"Number of variables: {len(batch_df.columns)}")

    table = doc.add_table(rows=1, cols=len(batch_df.columns))
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for i, col in enumerate(batch_df.columns):
        header_cells[i].text = str(col)

    for _, row in batch_df.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(batch_df.columns):
            cells[i].text = safe_text(row.get(col, ""))

    doc.save(path)

    return path

